from __future__ import annotations

import json
import os
import re
import zipfile
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

from invoice_tool.customer_reimbursements.export_attachments import exportable_settlement_attachments

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


EXPORT_LABELS = {
    "daily_reports": "日报",
    "mileage_proofs": "批导里程",
    "settlements": "工单结算",
    "invoices": "发票",
}


def safe_name(value, fallback="file"):
    text = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "-", str(value or "")).strip(" .")
    return text or fallback


def utcnow_text():
    return datetime.now(timezone.utc).isoformat()


def ensure_export_job_schema(db):
    """Fail clearly when migration 0317 has not been applied.

    PostgreSQLConnection deliberately rejects DDL from request code, so schema
    creation belongs to the migration step rather than the web process.
    """
    try:
        db().execute("select id from batch_export_jobs limit 1").fetchone()
    except Exception as exc:
        raise RuntimeError("batch_export_jobs schema is missing; apply PostgreSQL migration 0317") from exc


def _money(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def settlement_rows(reimbursement, items, expense_amount):
    """Return export rows; independent rental fuel is shown as a real detail row."""
    rows = []
    for index, raw in enumerate(items, 1):
        item = dict(raw)
        expenses = [float(expense_amount(item, key) or 0) for key in (
            "lodging", "airfare", "baggage", "rental_car", "fuel", "parking", "taxi"
        )]
        rows.append([
            index, item.get("worker_name") or "", item.get("project_date") or "",
            _money(item.get("standard_hours")), _money(item.get("transport_hours")),
            _money(item.get("public_transport_hours")), _money(item.get("overtime_hours")),
            _money(item.get("holiday_hours")), _money(item.get("labor_total")), *expenses,
            _money(item.get("mileage_total")), float(expense_amount(item, "other") or 0),
            _money(item.get("total")),
        ])
    rental_fuel = _money(reimbursement.get("rental_fuel_total") if hasattr(reimbursement, "get") else reimbursement["rental_fuel_total"])
    if rental_fuel:
        # Do not hide this amount in the snapshot total. Show it under the dedicated fuel column.
        rows.append([len(rows) + 1, "租车油费", "", 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, rental_fuel, 0, 0, 0, 0, rental_fuel])
    visible_total = sum(_money(row[18]) for row in rows)
    adjustment = round(_money(reimbursement["total_amount"]) - visible_total, 2)
    if abs(adjustment) >= 0.005:
        # Discounts or snapshot adjustments must also be visible so detail rows reconcile to the snapshot total.
        rows.append([len(rows) + 1, "结算折扣/调整", "", 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, adjustment])
    return rows


def build_settlement_xlsx(reimbursement, order, items, *, expense_amount, build_simple_xlsx):
    headers = ["序号", "姓名", "项目日期", "标准工时", "交通工时", "公共交通工时", "加班工时", "假期工时", "人工费", "住宿费", "机票费", "行李费", "租车费", "燃油费", "停车费", "出租车费", "里程费", "其他", "合计"]
    rows = settlement_rows(reimbursement, items, expense_amount)
    totals = ["合计", "", ""] + [0] * 16
    for row in rows:
        for idx in range(3, 19):
            totals[idx] += _money(row[idx])
    # Snapshot remains authoritative, but visible columns now reconcile because rental fuel is explicit.
    totals[18] = _money(reimbursement["total_amount"])
    rows.append(totals)
    rows.append(["说明：租车油费已在「燃油费」列单独展示；MRO 耗材已计入「其他」列；合计为工单结算快照总额。"] + [""] * 18)
    return build_simple_xlsx(headers, rows, sheet_name="工单结算")


def build_master_ledger(settlement_payloads, *, expense_amount):
    wb = Workbook()
    ws = wb.active
    ws.title = "工单结算"
    ws.merge_cells("A1:A2"); ws["A1"] = "派工单号"
    ws.merge_cells("B1:B2"); ws["B1"] = "序号"
    ws.merge_cells("C1:C2"); ws["C1"] = "姓名"
    ws.merge_cells("D1:D2"); ws["D1"] = "项目时间"
    ws.merge_cells("E1:M1"); ws["E1"] = "工时"
    ws.merge_cells("N1:X1"); ws["N1"] = "交通"
    ws.merge_cells("Y1:Y2"); ws["Y1"] = "合计"
    detail_headers = {
        "E2": "标准\n工时", "F2": "交通\n工时", "G2": "加班\n工时", "H2": "节假日\n工时",
        "I2": "标准工时费/小时", "J2": "交通工时费/小时", "K2": "加班工时费/小时", "L2": "节假日工时费/小时",
        "M2": "工时费\n合计", "N2": "住宿", "O2": "机票", "P2": "行李", "Q2": "租车", "R2": "加油",
        "S2": "停车", "T2": "打车", "U2": "公里数", "V2": "里程费/英里", "W2": "里程费", "X2": "其他",
    }
    for cell_ref, value in detail_headers.items():
        ws[cell_ref] = value
    header_fill = PatternFill("solid", fgColor="E2F0D9")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=1, max_row=2, min_col=1, max_col=25):
        for cell in row:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
    seq = 1
    excel_row = 3
    grand_total = 0.0
    for payload in settlement_payloads:
        reimbursement, order, items = payload
        order_start_row = ws.max_row + 1
        first = True
        for raw in items:
            item = dict(raw)
            values = [
                order["order_number"] if first else "", seq, item.get("worker_name") or "", item.get("project_date") or "",
                _money(item.get("standard_hours")), _money(item.get("transport_hours")), _money(item.get("overtime_hours")), _money(item.get("holiday_hours")),
                _money(item.get("standard_rate")), _money(item.get("transport_rate")), _money(item.get("overtime_rate")), _money(item.get("holiday_rate")), _money(item.get("labor_total")),
                float(expense_amount(item, "lodging") or 0), float(expense_amount(item, "airfare") or 0), float(expense_amount(item, "baggage") or 0),
                float(expense_amount(item, "rental_car") or 0), float(expense_amount(item, "fuel") or 0), float(expense_amount(item, "parking") or 0), float(expense_amount(item, "taxi") or 0),
                _money(item.get("miles")), _money(item.get("mileage_rate")), _money(item.get("mileage_total")), float(expense_amount(item, "other") or 0), _money(item.get("total")),
            ]
            ws.append(values)
            first = False; seq += 1; excel_row += 1
        rental_fuel = _money(reimbursement["rental_fuel_total"])
        visible_order_total = sum(_money(dict(item).get("total")) for item in items)
        if rental_fuel:
            values = [order["order_number"] if first else "", seq, "租车油费", "", 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, rental_fuel, 0, 0, 0, 0, 0, 0, rental_fuel]
            ws.append(values)
            first = False; seq += 1; excel_row += 1; visible_order_total += rental_fuel
        snapshot_total = _money(reimbursement["total_amount"])
        adjustment = round(snapshot_total - visible_order_total, 2)
        if abs(adjustment) >= 0.005:
            values = [order["order_number"] if first else "", seq, "结算折扣/调整", "", 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, adjustment]
            ws.append(values)
            first = False; seq += 1; excel_row += 1
        order_end_row = ws.max_row
        if order_end_row >= order_start_row:
            ws.cell(order_start_row, 1, order["order_number"])
            if order_end_row > order_start_row:
                ws.merge_cells(start_row=order_start_row, start_column=1, end_row=order_end_row, end_column=1)
            ws.cell(order_start_row, 1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        grand_total += snapshot_total
    total_row = ws.max_row + 1
    ws.cell(total_row, 1, "合计")
    ws.cell(total_row, 25, grand_total)
    note_row = total_row + 1
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=25)
    ws.cell(note_row, 1, "说明：租车油费已在「加油」列单独展示；MRO 耗材已计入「其他」列；如有折扣或快照调整，会以「结算折扣/调整」行展示；合计为各工单结算快照总额。")
    for row in ws.iter_rows(min_row=3, max_row=total_row, min_col=1, max_col=25):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.column >= 5 and cell.column <= 25:
                cell.number_format = "0.00"
    for cell in ws[total_row]:
        cell.font = Font(bold=True); cell.border = border
    ws.cell(note_row, 1).font = Font(italic=True, color="666666")
    widths = [20,8,18,14,10,10,10,10,15,15,15,15,15,12,12,12,12,12,12,12,12,14,12,12,14]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64+idx) if idx <= 26 else "A"].width = width
    ws.freeze_panes = "A3"
    buffer = BytesIO(); wb.save(buffer); buffer.seek(0)
    return buffer


def run_export_job(job_id, *, app, db, data_dir, build_service_report_docx, build_simple_xlsx,
                   customer_reimbursement_items, expense_amount, render_invoice_pdf, report_attachment_path):
    with app.app_context():
        ensure_export_job_schema(db)
        conn = db()
        job = conn.execute("select * from batch_export_jobs where id=?", (job_id,)).fetchone()
        if not job:
            return
        order_ids = [int(v) for v in json.loads(job["order_ids"] or "[]")]
        export_type = job["export_type"]
        conn.execute("update batch_export_jobs set status='processing', progress=1, started_at=? where id=?", (utcnow_text(), job_id)); conn.commit()
        try:
            root = Path(data_dir) / "batch-exports"; root.mkdir(parents=True, exist_ok=True)
            zip_name = f"{EXPORT_LABELS.get(export_type, export_type)}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.zip"
            stored = root / f"{job_id}.zip"
            missing = []
            settlement_payloads = []
            with zipfile.ZipFile(stored, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                total = max(1, len(order_ids))
                for pos, order_id in enumerate(order_ids, 1):
                    order = conn.execute(
                        """select service_orders.*,
                                  coalesce(owners.name, buyers.owner) as buyer_owner
                           from service_orders
                           left join buyers on buyers.id = service_orders.buyer_id
                           left join owners on owners.id = buyers.owner_id
                           where service_orders.id=?""",
                        (order_id,),
                    ).fetchone()
                    if not order:
                        missing.append(f"工单 {order_id}: 不存在")
                        continue
                    order_no = safe_name(order["order_number"], str(order_id))
                    if export_type == "daily_reports":
                        reports = conn.execute("select * from service_reports where service_order_id=? order by report_date,id", (order_id,)).fetchall()
                        if not reports:
                            missing.append(f"{order_no}: 无日报")
                        for report in reports:
                            data = build_service_report_docx(report, order)
                            name = safe_name(f"{order.get('client_order_number') if hasattr(order,'get') else order['client_order_number'] or order_no}-{report['report_date']}-report.docx", "service-report.docx")
                            archive.writestr(f"日报/{order_no}/{name}", data)
                    elif export_type == "mileage_proofs":
                        reports = conn.execute(
                            "select id, report_date from service_reports where service_order_id=? order by report_date,id",
                            (order_id,),
                        ).fetchall()
                        if not reports:
                            missing.append(f"{order_no}: 无日报")
                        for report in reports:
                            report_date = safe_name(report["report_date"] or f"report-{report['id']}", f"report-{report['id']}")
                            attachments = conn.execute(
                                "select * from service_report_attachments where report_id=? and category='mileage_proof' order by id",
                                (report["id"],),
                            ).fetchall()
                            if not attachments:
                                missing.append(f"{order_no} / {report_date}: 无里程佐证")
                                continue
                            for attachment in attachments:
                                source_path = report_attachment_path(attachment)
                                if not source_path or not os.path.isfile(source_path):
                                    missing.append(f"{order_no} / {report_date}: 里程佐证文件缺失 - {attachment['original_filename'] or attachment['id']}")
                                    continue
                                filename = safe_name(
                                    attachment["original_filename"] or os.path.basename(attachment["stored_filename"]),
                                    f"mileage-proof-{attachment['id']}",
                                )
                                archive.write(source_path, f"批导里程/{order_no}/{report_date}/{filename}")
                    elif export_type == "settlements":
                        reimbursement = conn.execute("select * from customer_reimbursements where service_order_id=? and is_current=true order by id desc limit 1", (order_id,)).fetchone()
                        if not reimbursement:
                            missing.append(f"{order_no}: 无工单结算")
                        else:
                            items = customer_reimbursement_items(reimbursement["id"])
                            settlement_payloads.append((reimbursement, order, items))
                            xlsx = build_settlement_xlsx(reimbursement, order, items, expense_amount=expense_amount, build_simple_xlsx=build_simple_xlsx)
                            data = xlsx.getvalue() if hasattr(xlsx, "getvalue") else xlsx.read()
                            archive.writestr(f"工单结算/{order_no}_工单结算.xlsx", data)
                            attachment_root = Path(data_dir) / "customer-reimbursements" / str(reimbursement["id"]) / "attachments"
                            attachment_names = set()
                            attachment_sources = set()
                            for attachment in exportable_settlement_attachments(conn, reimbursement["id"]):
                                source_path = attachment_root / attachment["stored_filename"]
                                source_path = source_path.resolve()
                                if not source_path.is_relative_to(attachment_root.resolve()):
                                    raise ValueError("Invalid settlement attachment path")
                                if source_path in attachment_sources:
                                    continue
                                attachment_sources.add(source_path)
                                if not source_path.is_file():
                                    missing.append(
                                        f"{order_no}: 结算附件文件缺失 - {attachment['original_filename'] or attachment['id']}"
                                    )
                                    continue
                                filename = safe_name(
                                    attachment["original_filename"] or attachment["stored_filename"],
                                    f"attachment-{attachment['id']}",
                                )
                                stem, suffix = os.path.splitext(filename)
                                candidate, index = filename, 2
                                while candidate.casefold() in attachment_names:
                                    candidate = f"{stem}-{index}{suffix}"
                                    index += 1
                                attachment_names.add(candidate.casefold())
                                filename = candidate
                                archive.write(source_path, f"工单结算/{order_no}_附件/{filename}")
                    elif export_type == "invoices":
                        invoices = conn.execute("select * from invoices where service_order_id=? and status!='void' order by id", (order_id,)).fetchall()
                        if not invoices:
                            missing.append(f"{order_no}: 无发票")
                        for invoice in invoices:
                            client = conn.execute("select * from clients where id=?", (invoice["client_id"],)).fetchone()
                            items = conn.execute("select * from invoice_items where invoice_id=? order by id", (invoice["id"],)).fetchall()
                            pdf = render_invoice_pdf(invoice, client, items)
                            pdf_bytes = pdf.getvalue() if hasattr(pdf, "getvalue") else bytes(pdf)
                            archive.writestr(f"发票/{safe_name(invoice['invoice_number'], str(invoice['id']))}.pdf", pdf_bytes)
                    conn.execute("update batch_export_jobs set progress=? where id=?", (int(pos * 90 / total), job_id)); conn.commit()
                if export_type == "settlements" and settlement_payloads:
                    ledger = build_master_ledger(settlement_payloads, expense_amount=expense_amount)
                    archive.writestr("工单结算/工单结算总账单.xlsx", ledger.getvalue())
                if missing:
                    archive.writestr("缺失说明.txt", ("\n".join(missing) + "\n").encode("utf-8"))
            conn.execute("update batch_export_jobs set status='completed', progress=100, file_name=?, stored_path=?, completed_at=? where id=?", (zip_name, str(stored), utcnow_text(), job_id)); conn.commit()
        except Exception as exc:
            app.logger.exception("batch export job failed: %s", job_id)
            conn.execute("update batch_export_jobs set status='failed', error=?, completed_at=? where id=?", (str(exc)[:2000], utcnow_text(), job_id)); conn.commit()
