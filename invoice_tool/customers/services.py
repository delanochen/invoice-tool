"""Customer, owner, manufacturer, and site domain helpers."""

import csv
import re
from io import BytesIO

from flask import request
from werkzeug.utils import secure_filename


BUYER_IMPORT_FIELD_ALIASES = {
    "buyer_number": {"编号", "站点编号", "客户编号", "buyer_number", "site_number", "number"},
    "name": {"站点名称", "站点", "名称", "name", "site", "site_name"},
    "owner": {"业主", "业主名称", "owner", "owner_name"},
    "country_code": {"国家", "国家代码", "country", "country_code"},
    "contact_name": {"联系人", "contact", "contact_name"},
    "contact_details": {"联系方式", "联系电话", "电话", "contact_details", "phone", "tel"},
    "email": {"电子邮箱", "电子邮箱地址", "邮箱", "email"},
    "site_size": {"规模", "site_size", "scale", "size"},
    "equipment_manufacturer": {"配套机厂家", "配套厂家", "厂家", "equipment_manufacturer", "manufacturer"},
    "detailed_address": {"详细地址", "地址", "站点地址", "detailed_address", "address", "site_address"},
}

# 美国 50 州 + 华盛顿特区 + 常用领地简写。站点州简写只认这份清单，避免把街道后缀
# （ST/RD/AVE/NW 等）或非美国地址（ON/BC/QC 等）误当作州。
US_STATE_CODES = {
    "AL", "AK", "AZ", "AR", "CA", "CO", "CT", "DE", "FL", "GA",
    "HI", "ID", "IL", "IN", "IA", "KS", "KY", "LA", "ME", "MD",
    "MA", "MI", "MN", "MS", "MO", "MT", "NE", "NV", "NH", "NJ",
    "NM", "NY", "NC", "ND", "OH", "OK", "OR", "PA", "RI", "SC",
    "SD", "TN", "TX", "UT", "VT", "VA", "WA", "WV", "WI", "WY",
    "DC", "PR", "VI", "GU", "AS", "MP",
}

_STATE_TAIL_RE = re.compile(r"(?:^|[\s,])([A-Z]{2})\s*(?:[,.]?\s*\d{5}(?:-\d{4})?)?\s*$")
_STANDALONE_CODE_RE = re.compile(r"(?<![A-Z])([A-Z]{2})(?![A-Z])")


def state_code_from_address(address):
    """从美国地址文本解析州简写（如 "Houston, TX 77001" → "TX"）。

    优先匹配地址尾部的 "州简写 + 邮编" 或独立州简写；解析不到返回 None。
    只接受 US_STATE_CODES 清单内的简写，街道后缀（ST/RD/NW）与非美国地址
    （Toronto, ON / Vancouver, BC）不会被误判。
    """
    text = str(address or "").strip()
    if not text:
        return None
    upper = re.sub(r"\s+", " ", text.upper()).strip()
    match = _STATE_TAIL_RE.search(upper)
    if match and match.group(1) in US_STATE_CODES:
        return match.group(1)
    for code in _STANDALONE_CODE_RE.findall(upper):
        if code in US_STATE_CODES:
            return code
    return None


# 美国州/领地简写 → 各语言州名。en 为英文名；zh 为中文译名；es 为西班牙语名；
# nl/de 通常直接使用英文州名（语言代码未命中时统一回退 en）。
US_STATE_NAMES = {
    "AL": {"en": "Alabama", "zh": "阿拉巴马州", "es": "Alabama"},
    "AK": {"en": "Alaska", "zh": "阿拉斯加州", "es": "Alaska"},
    "AZ": {"en": "Arizona", "zh": "亚利桑那州", "es": "Arizona"},
    "AR": {"en": "Arkansas", "zh": "阿肯色州", "es": "Arkansas"},
    "CA": {"en": "California", "zh": "加利福尼亚州", "es": "California"},
    "CO": {"en": "Colorado", "zh": "科罗拉多州", "es": "Colorado"},
    "CT": {"en": "Connecticut", "zh": "康涅狄格州", "es": "Connecticut"},
    "DE": {"en": "Delaware", "zh": "特拉华州", "es": "Delaware"},
    "FL": {"en": "Florida", "zh": "佛罗里达州", "es": "Florida"},
    "GA": {"en": "Georgia", "zh": "佐治亚州", "es": "Georgia"},
    "HI": {"en": "Hawaii", "zh": "夏威夷州", "es": "Hawái"},
    "ID": {"en": "Idaho", "zh": "爱达荷州", "es": "Idaho"},
    "IL": {"en": "Illinois", "zh": "伊利诺伊州", "es": "Illinois"},
    "IN": {"en": "Indiana", "zh": "印第安纳州", "es": "Indiana"},
    "IA": {"en": "Iowa", "zh": "艾奥瓦州", "es": "Iowa"},
    "KS": {"en": "Kansas", "zh": "堪萨斯州", "es": "Kansas"},
    "KY": {"en": "Kentucky", "zh": "肯塔基州", "es": "Kentucky"},
    "LA": {"en": "Louisiana", "zh": "路易斯安那州", "es": "Luisiana"},
    "ME": {"en": "Maine", "zh": "缅因州", "es": "Maine"},
    "MD": {"en": "Maryland", "zh": "马里兰州", "es": "Maryland"},
    "MA": {"en": "Massachusetts", "zh": "马萨诸塞州", "es": "Massachusetts"},
    "MI": {"en": "Michigan", "zh": "密歇根州", "es": "Míchigan"},
    "MN": {"en": "Minnesota", "zh": "明尼苏达州", "es": "Minnesota"},
    "MS": {"en": "Mississippi", "zh": "密西西比州", "es": "Misisipi"},
    "MO": {"en": "Missouri", "zh": "密苏里州", "es": "Misuri"},
    "MT": {"en": "Montana", "zh": "蒙大拿州", "es": "Montana"},
    "NE": {"en": "Nebraska", "zh": "内布拉斯加州", "es": "Nebraska"},
    "NV": {"en": "Nevada", "zh": "内华达州", "es": "Nevada"},
    "NH": {"en": "New Hampshire", "zh": "新罕布什尔州", "es": "Nuevo Hampshire"},
    "NJ": {"en": "New Jersey", "zh": "新泽西州", "es": "Nueva Jersey"},
    "NM": {"en": "New Mexico", "zh": "新墨西哥州", "es": "Nuevo México"},
    "NY": {"en": "New York", "zh": "纽约州", "es": "Nueva York"},
    "NC": {"en": "North Carolina", "zh": "北卡罗来纳州", "es": "Carolina del Norte"},
    "ND": {"en": "North Dakota", "zh": "北达科他州", "es": "Dakota del Norte"},
    "OH": {"en": "Ohio", "zh": "俄亥俄州", "es": "Ohio"},
    "OK": {"en": "Oklahoma", "zh": "俄克拉何马州", "es": "Oklahoma"},
    "OR": {"en": "Oregon", "zh": "俄勒冈州", "es": "Oregón"},
    "PA": {"en": "Pennsylvania", "zh": "宾夕法尼亚州", "es": "Pensilvania"},
    "RI": {"en": "Rhode Island", "zh": "罗得岛州", "es": "Rhode Island"},
    "SC": {"en": "South Carolina", "zh": "南卡罗来纳州", "es": "Carolina del Sur"},
    "SD": {"en": "South Dakota", "zh": "南达科他州", "es": "Dakota del Sur"},
    "TN": {"en": "Tennessee", "zh": "田纳西州", "es": "Tennessee"},
    "TX": {"en": "Texas", "zh": "得克萨斯州", "es": "Texas"},
    "UT": {"en": "Utah", "zh": "犹他州", "es": "Utah"},
    "VT": {"en": "Vermont", "zh": "佛蒙特州", "es": "Vermont"},
    "VA": {"en": "Virginia", "zh": "弗吉尼亚州", "es": "Virginia"},
    "WA": {"en": "Washington", "zh": "华盛顿州", "es": "Washington"},
    "WV": {"en": "West Virginia", "zh": "西弗吉尼亚州", "es": "Virginia Occidental"},
    "WI": {"en": "Wisconsin", "zh": "威斯康星州", "es": "Wisconsin"},
    "WY": {"en": "Wyoming", "zh": "怀俄明州", "es": "Wyoming"},
    "DC": {"en": "District of Columbia", "zh": "华盛顿哥伦比亚特区", "es": "Distrito de Columbia"},
    "PR": {"en": "Puerto Rico", "zh": "波多黎各", "es": "Puerto Rico"},
    "VI": {"en": "U.S. Virgin Islands", "zh": "美属维尔京群岛", "es": "Islas Vírgenes de los Estados Unidos"},
    "GU": {"en": "Guam", "zh": "关岛", "es": "Guam"},
    "AS": {"en": "American Samoa", "zh": "美属萨摩亚", "es": "Samoa Americana"},
    "MP": {"en": "Northern Mariana Islands", "zh": "北马里亚纳群岛", "es": "Islas Marianas del Norte"},
}


def _normalize_state_language(language):
    lang = (language or "zh-CN").lower()
    return "zh" if lang == "zh-cn" else lang


def localized_state_name(state_code, language="zh-CN"):
    """返回指定语言的州名（如 "TX" + "zh-CN" → "得克萨斯州"）；无此州或简写为空返回 ""。"""
    names = US_STATE_NAMES.get((state_code or "").upper(), {})
    if not names:
        return ""
    lang = _normalize_state_language(language)
    return names.get(lang, names.get("en", ""))


def state_names_for_language(language="zh-CN"):
    """返回 州简写 → 本地化州名 的映射（供表单 JS 在地址变化时联动更新州名）。"""
    lang = _normalize_state_language(language)
    return {
        code: names.get(lang, names.get("en", ""))
        for code, names in US_STATE_NAMES.items()
    }


def build_customer_services(
    *, db, lock_number_allocation, now, normalized_address, country_by_code
):
    def next_client_number():
        lock_number_allocation(db())
        row = db().execute(
            """
            select client_number from clients
            where client_number glob '[0-9][0-9][0-9][0-9][0-9]'
            order by client_number desc limit 1
            """
        ).fetchone()
        return "00001" if not row else f"{int(row['client_number']) + 1:05d}"

    def next_buyer_number():
        lock_number_allocation(db())
        row = db().execute(
            """
            select buyer_number from buyers
            where buyer_number glob 'BUY[0-9][0-9][0-9][0-9][0-9]'
            order by buyer_number desc limit 1
            """
        ).fetchone()
        return "BUY00001" if not row else f"BUY{int(row['buyer_number'][3:]) + 1:05d}"

    def next_owner_number():
        lock_number_allocation(db())
        row = db().execute(
            """
            select owner_number from owners
            where owner_number glob 'OWN[0-9][0-9][0-9][0-9][0-9]'
            order by owner_number desc limit 1
            """
        ).fetchone()
        return "OWN00001" if not row else f"OWN{int(row['owner_number'][3:]) + 1:05d}"

    def next_manufacturer_number():
        lock_number_allocation(db())
        row = db().execute(
            """
            select manufacturer_number from manufacturers
            where manufacturer_number glob 'MFG[0-9][0-9][0-9][0-9][0-9]'
            order by manufacturer_number desc limit 1
            """
        ).fetchone()
        return "MFG00001" if not row else f"MFG{int(row['manufacturer_number'][3:]) + 1:05d}"

    def unknown_owner():
        owner = db().execute("select * from owners where name = ?", ("未知",)).fetchone()
        if owner:
            return owner
        cursor = db().execute(
            "insert into owners (owner_number, name, created_at) values (?, ?, ?)",
            (next_owner_number(), "未知", now()),
        )
        return db().execute("select * from owners where id = ?", (cursor.lastrowid,)).fetchone()

    def owner_options():
        return db().execute(
            """
            select id, owner_number, name from owners
            order by case when name = '未知' then 0 else 1 end, owner_number
            """
        ).fetchall()

    def manufacturer_options():
        return db().execute(
            """
            select id, manufacturer_number, name from manufacturers
            order by manufacturer_number
            """
        ).fetchall()

    def manufacturer_from_form():
        manufacturer_id = (
            int(request.form["manufacturer_id"])
            if request.form.get("manufacturer_id", "").isdigit()
            else None
        )
        if not manufacturer_id:
            return None
        return db().execute("select * from manufacturers where id = ?", (manufacturer_id,)).fetchone()

    def manufacturer_by_name(name):
        clean_name = " ".join(str(name or "").strip().split())
        if not clean_name:
            return None
        return db().execute(
            "select * from manufacturers where lower(trim(name)) = lower(trim(?))",
            (clean_name,),
        ).fetchone()

    def normalize_import_header(value):
        return re.sub(r"[\s_（）()：:.-]+", "", str(value or "").strip()).casefold()

    BUYER_IMPORT_HEADER_LOOKUP = {
        normalize_import_header(alias): field
        for field, aliases in BUYER_IMPORT_FIELD_ALIASES.items()
        for alias in aliases
    }

    def get_or_create_owner_by_name(owner_name):
        name = " ".join(str(owner_name or "").strip().split()) or "未知"
        row = db().execute("select * from owners where lower(trim(name)) = lower(trim(?))", (name,)).fetchone()
        if row:
            return row, False
        cursor = db().execute(
            "insert into owners (owner_number, name, created_at) values (?, ?, ?)",
            (next_owner_number(), name, now()),
        )
        return db().execute("select * from owners where id = ?", (cursor.lastrowid,)).fetchone(), True

    def country_from_import_value(value, default_code="US"):
        text = str(value or "").strip()
        if text:
            country = country_by_code(text, include_inactive=True)
            if country:
                return country
            country = db().execute(
                """
                select countries.*
                from countries
                left join country_translations on country_translations.country_code = countries.code
                where lower(country_translations.name) = lower(?)
                limit 1
                """,
                (text,),
            ).fetchone()
            if country:
                return country
        return country_by_code(default_code, include_inactive=True)

    def imported_buyer_rows(uploaded_file):
        filename = secure_filename(uploaded_file.filename or "")
        extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        content = uploaded_file.read()
        if not content:
            raise ValueError("导入文件为空。")
        if extension == "csv":
            for encoding in ("utf-8-sig", "gb18030"):
                try:
                    text = content.decode(encoding)
                    break
                except UnicodeDecodeError:
                    text = ""
            if not text:
                raise ValueError("CSV 文件编码无法识别，请使用 UTF-8 或 GB18030。")
            rows = list(csv.DictReader(text.splitlines()))
        elif extension == "xlsx":
            try:
                from openpyxl import load_workbook
            except ImportError as error:
                raise ValueError("服务器未安装 Excel 导入组件，请先安装 openpyxl。") from error
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            worksheet = workbook.active
            values = list(worksheet.iter_rows(values_only=True))
            if not values:
                return []
            headers = [str(cell or "").strip() for cell in values[0]]
            rows = [
                {headers[index]: cell for index, cell in enumerate(row) if index < len(headers)}
                for row in values[1:]
            ]
        else:
            raise ValueError("仅支持导入 .xlsx 或 .csv 文件。")
        normalized_rows = []
        for source_row in rows:
            normalized = {}
            for header, value in source_row.items():
                field = BUYER_IMPORT_HEADER_LOOKUP.get(normalize_import_header(header))
                if field:
                    normalized[field] = "" if value is None else str(value).strip()
            if any(normalized.values()):
                normalized_rows.append(normalized)
        return normalized_rows

    def import_buyers_from_file(uploaded_file, client_id=None):
        rows = imported_buyer_rows(uploaded_file)
        result = {
            "total": len(rows),
            "created": 0,
            "skipped": 0,
            "errors": 0,
            "owners_created": 0,
            "details": [],
        }
        seen_keys = set()
        for index, row in enumerate(rows, start=2):
            name = row.get("name", "").strip()
            detailed_address = row.get("detailed_address", "").strip()
            buyer_number = row.get("buyer_number", "").strip().upper()
            if not name or not detailed_address:
                result["errors"] += 1
                result["details"].append(f"第 {index} 行：缺少站点名称或详细地址，未导入。")
                continue
            duplicate_key = (client_id or 0, normalized_address(name), normalized_address(detailed_address))
            if duplicate_key in seen_keys:
                result["skipped"] += 1
                result["details"].append(f"第 {index} 行：{name} 在本次文件中重复，已跳过。")
                continue
            seen_keys.add(duplicate_key)
            existing_site = db().execute(
                """
                select buyer_number from buyers
                where coalesce(client_id, 0) = coalesce(?, 0)
                  and lower(trim(name)) = lower(trim(?))
                  and lower(trim(detailed_address)) = lower(trim(?))
                """,
                (client_id, name, detailed_address),
            ).fetchone()
            if existing_site:
                result["skipped"] += 1
                result["details"].append(f"第 {index} 行：{name} 已存在（{existing_site['buyer_number']}），已跳过。")
                continue
            if buyer_number and db().execute("select id from buyers where buyer_number = ?", (buyer_number,)).fetchone():
                result["skipped"] += 1
                result["details"].append(f"第 {index} 行：编号 {buyer_number} 已存在，已跳过。")
                continue
            owner, owner_created = get_or_create_owner_by_name(row.get("owner"))
            country = country_from_import_value(row.get("country_code"), "US")
            manufacturer = manufacturer_by_name(row.get("equipment_manufacturer"))
            db().execute(
                """
                insert into buyers (
                    buyer_number, client_id, country, country_code, name, owner_id, owner, manufacturer_id, contact_name, contact_details,
                    email, site_size, detailed_address, equipment_manufacturer, state_code, created_at
                ) values (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    buyer_number or next_buyer_number(),
                    client_id,
                    country["code"],
                    country["code"],
                    name,
                    owner["id"],
                    owner["name"],
                    manufacturer["id"] if manufacturer else None,
                    row.get("contact_name", ""),
                    row.get("contact_details", ""),
                    row.get("email", ""),
                    row.get("site_size", ""),
                    detailed_address,
                    row.get("equipment_manufacturer", ""),
                    state_code_from_address(detailed_address) or "",
                    now(),
                ),
            )
            result["created"] += 1
            if owner_created:
                result["owners_created"] += 1
        return result

    return {
        "next_client_number": next_client_number,
        "next_buyer_number": next_buyer_number,
        "next_owner_number": next_owner_number,
        "next_manufacturer_number": next_manufacturer_number,
        "unknown_owner": unknown_owner,
        "owner_options": owner_options,
        "manufacturer_options": manufacturer_options,
        "manufacturer_from_form": manufacturer_from_form,
        "manufacturer_by_name": manufacturer_by_name,
        "normalize_import_header": normalize_import_header,
        "get_or_create_owner_by_name": get_or_create_owner_by_name,
        "country_from_import_value": country_from_import_value,
        "imported_buyer_rows": imported_buyer_rows,
        "import_buyers_from_file": import_buyers_from_file,
    }
