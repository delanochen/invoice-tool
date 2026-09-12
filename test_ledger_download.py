import unittest
from io import BytesIO
from zipfile import ZipFile
from test_field_work import FieldWorkTest


class LedgerDownloadTest(unittest.TestCase):
    def setUp(self):
        self.f = FieldWorkTest()
        self.f.setUp()
        self.addCleanup(self.f.doCleanups)
        self.assertEqual(self.f.upload(equipment_number='M-101', position_number='P-7', container_number='C-8').status_code, 200)
        self.assertEqual(self.f.upload(equipment_number='M-202', position_number='P-7', container_number='C-9').status_code, 200)

    def test_filters_and_zip_use_identical_rows(self):
        filters = dict(equipment_number='101', position_number='p-7', container_number='C-8')
        rows = self.f.http.get('/api/field/photos', query_string=filters).json['rows']
        self.assertEqual(len(rows), 1)
        response = self.f.http.get('/api/field/photos.zip', query_string=filters)
        self.assertEqual(response.status_code, 200)
        with ZipFile(BytesIO(response.data)) as archive:
            self.assertEqual(len(archive.namelist()), 1)
            name = archive.namelist()[0]
            self.assertIn('/'+str(rows[0]['id'])+'-', name)
            self.assertGreater(len(archive.read(name)), 0)
        response.close()
        page = self.f.http.get('/reports/field-photos', query_string=filters)
        self.assertEqual(page.status_code, 200)
        for field in filters:
            self.assertIn(f'name="{field}"', page.text)
        self.assertEqual(self.f.http.get('/api/field/photos', query_string=dict(container_number='C-9', equipment_number='101')).json['rows'], [])

    def test_excel_export_contains_photo_thumbnails_and_split_customer_site(self):
        from openpyxl import load_workbook

        response = self.f.http.get('/api/field/photos.xlsx')
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.data))
        sheet = workbook.active
        self.assertEqual(sheet.cell(1, 1).value, '照片')
        self.assertEqual(sheet.cell(1, 3).value, '客户')
        self.assertEqual(sheet.cell(1, 4).value, '站点')
        self.assertEqual(len(sheet._images), 2)
        workbook.close()

    def test_empty_zip_and_unauthenticated_access(self):
        self.assertEqual(self.f.http.get('/api/field/photos.zip?equipment_number=no-match').status_code, 422)
        with self.f.http.session_transaction() as session:
            session.clear()
        self.assertEqual(self.f.http.get('/api/field/photos.zip').status_code, 401)

    def test_position_and_container_normalized_on_upload_and_existing_records(self):
        import field_work
        response = self.f.upload(position_number='  ab-1\t', container_number='\n box-2  ')
        self.assertEqual(response.status_code, 200)
        with self.f.module.app.app_context():
            db = self.f.module.db()
            row = db.execute('select * from field_photos where id = ?', (response.json['id'],)).fetchone()
            self.assertEqual(row['position_number'], 'AB-1')
            self.assertEqual(row['container_number'], 'BOX-2')
            db.execute("update field_photos set position_number='  ab-1  ', container_number='  box-2  ' where id=?", (row['id'],))
            field_work.init_field_schema(db)
            saved = db.execute('select * from field_photos where id = ?', (row['id'],)).fetchone()
            self.assertEqual(saved['position_number'], 'AB-1')
            self.assertEqual(saved['container_number'], 'BOX-2')

    def test_repair_filters_match_export_and_technician_not_uploader(self):
        from openpyxl import load_workbook
        with self.f.module.app.app_context():
            db = self.f.module.db()
            db.execute("update field_photos set photo_type='equipment', technician_name='Repair Alice' where equipment_number='M-101'")
            db.execute("update field_photos set photo_type='equipment', technician_name='Repair Bob' where equipment_number='M-202'")
            db.commit()
        filters = dict(order_id=str(self.f.fixture.order), technician='alice', equipment_number='101', position_number='P-7', container_number='C-8')
        response = self.f.http.get('/reports/field-repairs', query_string=filters)
        self.assertEqual(response.status_code, 200)
        for field in filters:
            self.assertIn(f'name="{field}"', response.text)
        self.assertIn('<select name="order_id">', response.text)
        self.assertIn(f'value="{self.f.fixture.order}" selected', response.text)
        self.assertIn('Repair Alice', response.text)
        self.assertNotIn('Repair Bob', response.text)
        export = self.f.http.get('/api/field/repairs.xlsx', query_string=filters)
        self.assertEqual(export.status_code, 200)
        workbook = load_workbook(BytesIO(export.data), read_only=True)
        rows = list(workbook.active.values)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][4], 'M-101')
        self.assertEqual(rows[1][7], 'Repair Alice')
        workbook.close()
        no_order = self.f.http.get('/reports/field-repairs', query_string={**filters, 'order_number': 'NO-SUCH-ORDER'})
        self.assertIn('没有符合条件的设备照片。', no_order.text)
        filters['technician'] = 'Bob'
        empty = self.f.http.get('/reports/field-repairs', query_string=filters)
        self.assertIn('没有符合条件的设备照片。', empty.text)
