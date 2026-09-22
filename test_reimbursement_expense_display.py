"""工单结算单：报销来源金额的显示、已调整徽章、以及保存不被冲掉。

覆盖 v0.1.258 修复：
  1. 单元格显示「实际生效金额」——未调整时等于报销来源合计，而不是 0。
  2. 「已调整」徽章只在用户真的改过金额时出现，未改过时不出现。
  3. 用户改过的金额在保存（含 totals 重算、重进页面）后保持不变。
"""
import importlib.util
import json
import shutil
import tempfile
import unittest
from pathlib import Path


REPO_DIR = Path(__file__).resolve().parent
MONEY_FIELDS = ("lodging", "airfare", "baggage", "rental_car", "fuel", "parking", "taxi", "other")


class ReimbursementExpenseDisplayTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp_dir = tempfile.TemporaryDirectory()
        module_path = Path(cls.temp_dir.name) / "app.py"
        shutil.copyfile(REPO_DIR / "app.py", module_path)
        spec = importlib.util.spec_from_file_location("invoice_tool_reimb_display_test_app", module_path)
        cls.module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.module)
        cls.module.app.config.update(TESTING=True, SECRET_KEY="test-secret")

    @classmethod
    def tearDownClass(cls):
        cls.temp_dir.cleanup()

    def setUp(self):
        with self.module.app.app_context():
            connection = self.module.db()
            connection.execute("delete from customer_reimbursement_expense_links")
            connection.execute("delete from customer_reimbursement_items")
            connection.execute("delete from customer_reimbursements")
            connection.execute("delete from expense_items")
            connection.execute("delete from expenses")
            connection.execute("delete from service_orders where order_number = 'SO-DISPLAY'")
            connection.execute("delete from users where email = 'display-admin@example.com'")
            connection.execute("delete from projects where project_type = 'expense'")
            user_id = connection.execute(
                """
                insert into users (name, email, password_hash, role, created_at)
                values ('Display Admin', 'display-admin@example.com', 'unused', 'admin', '2026-09-15T12:00:00')
                """
            ).lastrowid
            order_id = connection.execute(
                """
                insert into service_orders (
                    order_number, client_name, site_address, client_order_number, start_date,
                    created_by, created_at
                ) values ('SO-DISPLAY', 'Test client', 'Test address', 'CLIENT-DISPLAY', '2026-09-15', ?, '2026-09-15T12:00:00')
                """,
                (user_id,),
            ).lastrowid
            lodging_project_id = connection.execute(
                """
                insert into projects (name, name_key, project_type, unit_price, is_active, created_at)
                values ('Accommodation/Lodging', 'accommodation/lodging', 'expense', 0, 1, '2026-09-15T12:00:00')
                """
            ).lastrowid
            expense_id = connection.execute(
                """
                insert into expenses (
                    service_order_id, expense_number, project, expense_date, amount, status,
                    created_by, beneficiary_id, created_at, updated_at
                ) values (?, 'EXP-LODGE', 'Accommodation/Lodging', '2026-09-15', 300, 'approved', ?, ?,
                          '2026-09-15T12:00:00', '2026-09-15T12:00:00')
                """,
                (order_id, user_id, user_id),
            ).lastrowid
            connection.execute(
                """
                insert into expense_items (expense_id, project_id, project, amount, sort_order)
                values (?, ?, 'Accommodation/Lodging', 300, 0)
                """,
                (expense_id, lodging_project_id),
            )
            reimbursement_id = connection.execute(
                """
                insert into customer_reimbursements (
                    service_order_id, file_name, stored_filename, status, expense_selection_mode,
                    created_by, created_at
                ) values (?, 'settlement.pdf', 'settlement.pdf', 'draft', 'legacy', ?, '2026-09-15T12:00:00')
                """,
                (order_id, user_id),
            ).lastrowid
            # 与该报销同姓名同日期的一行结算明细，来源金额应落到 lodging
            connection.execute(
                """
                insert into customer_reimbursement_items (
                    customer_reimbursement_id, worker_name, project_date,
                    standard_hours, standard_rate, labor_total, total, sort_order
                ) values (?, 'Display Admin', '2026-09-15', 0, 0, 0, 0, 0)
                """,
                (reimbursement_id,),
            )
            connection.commit()
            self.order_id = order_id
            self.user_id = user_id
            self.reimbursement_id = reimbursement_id

    def _item(self):
        with self.module.app.app_context():
            return dict(self.module.customer_reimbursement_items(self.reimbursement_id)[0])

    def _effective(self, row, field_name):
        with self.module.app.app_context():
            return float(self.module.customer_reimbursement_item_expense_amount(row, field_name))

    def _recalc(self, rows=None):
        """重算结算合计并提交。

        每个 app context 拿到的是独立 SQLite 连接，测试读取时必须 commit，
        否则后续 context 看不到这次写入。
        """
        with self.module.app.app_context():
            totals = self.module.update_customer_reimbursement_totals(self.reimbursement_id, rows)
            self.module.db().commit()
            return totals

    # ---- 1. 未调整时，生效金额等于报销来源金额，不是 0 ----
    def test_untouched_cell_uses_auto_transferred_amount(self):
        self._recalc()
        item = self._item()
        self.assertEqual(item["auto_lodging"], 300)
        self.assertEqual(item["lodging"], 0, "未调整时人工列为 0，代表沿用来源金额")
        self.assertEqual(self._effective(item, "lodging"), 300, "单元格生效金额应为来源金额")

    # ---- 2. 未调整时不应出现「已调整」徽章 ----
    def test_untouched_cell_has_no_adjusted_badge(self):
        self._recalc()
        item = self._item()
        # 模板判定：manual_amount and manual_amount != auto_amount
        manual = item["lodging"] or 0
        is_adjusted = bool(manual) and manual != item["auto_lodging"]
        self.assertFalse(is_adjusted, "未调整的单元格不应显示已调整")

    # ---- 3. 手改后徽章出现 ----
    def test_edited_cell_shows_adjusted_badge(self):
        self._recalc()
        item = self._item()
        item["lodging"] = 250.0
        manual = item["lodging"] or 0
        is_adjusted = bool(manual) and manual != item["auto_lodging"]
        self.assertTrue(is_adjusted)

    # ---- 4. 关键回归：手改值保存后不被来源金额冲掉 ----
    def test_manual_override_survives_save_and_recalc(self):
        self._recalc()
        item = self._item()
        item["lodging"] = 250.0           # 用户手改
        item.pop("auto_expense_sources", None)
        self._recalc([item])
        after = self._item()
        self.assertEqual(after["lodging"], 250.0, "用户手改的 250 必须保留")
        self.assertEqual(after["auto_lodging"], 300, "来源金额仍应记录为 300")
        self.assertEqual(
            self._effective(after, "lodging"),
            250.0,
            "生效金额应为手改后的 250，不能与来源相加",
        )

    # ---- 5. 表单回传 auto_* 后，totals 重算仍保留手改值 ----
    def test_form_roundtrip_preserves_manual_override(self):
        self._recalc()
        with self.module.app.test_client() as client:
            with client.session_transaction() as session:
                session["user_id"] = self.user_id
            with self.module.app.app_context():
                db = self.module.db()
                db.execute("update customer_reimbursement_items set lodging = 250 where customer_reimbursement_id = ?",
                           (self.reimbursement_id,))
                db.commit()
            # 保存（POST）→ 后端 from_form 读到 lodging=250 与 echo 回来的 auto_lodging=300
            response = client.post(
                f"/service-orders/{self.order_id}/customer-reimbursement",
                data={
                    "action": "save",
                    "worker_name": ["Display Admin"],
                    "project_date": ["2026-09-15"],
                    "standard_hours": ["0"],
                    "transport_hours": ["0"],
                    "public_transport_hours": ["0"],
                    "overtime_hours": ["0"],
                    "holiday_hours": ["0"],
                    "lodging": ["250"],
                    "auto_lodging": ["300"],
                    "airfare": ["0"], "auto_airfare": ["0"],
                    "baggage": ["0"], "auto_baggage": ["0"],
                    "rental_car": ["0"], "auto_rental_car": ["0"],
                    "fuel": ["0"], "auto_fuel": ["0"],
                    "parking": ["0"], "auto_parking": ["0"],
                    "taxi": ["0"], "auto_taxi": ["0"],
                    "miles": ["0"],
                    "other": ["0"], "auto_other": ["0"],
                    "source_report_id": [""],
                    "source_worker_user_id": [""],
                },
                follow_redirects=False,
            )
            self.assertIn(response.status_code, (302, 200))
            after = self._item()
            self.assertEqual(after["lodging"], 250, "保存后手改值不应被冲掉")
            self.assertEqual(after["auto_lodging"], 300)

    # ---- 6. 结算合计只计一次，不重复累加来源 ----
    def test_totals_do_not_double_count_sources(self):
        totals = self._recalc()
        self.assertEqual(totals["lodging_total"], 300, "来源 300 只能计一次")
        self.assertEqual(totals["total_amount"], 300)
        with self.module.app.app_context():
            again = self.module.customer_reimbursement_totals(
                self.module.customer_reimbursement_items(self.reimbursement_id)
            )
            self.assertEqual(again["lodging_total"], 300, "重复读取不应改变结果")
            self.assertEqual(again["total_amount"], 300)

    # ---- 7. 未调整 + 有来源：页面渲染的 input value 即为来源金额 ----
    def test_rendered_input_shows_source_amount(self):
        self._recalc()
        with self.module.app.test_client() as client:
            with client.session_transaction() as session:
                session["user_id"] = self.user_id
            response = client.get(f"/service-orders/{self.order_id}/customer-reimbursement")
            self.assertEqual(response.status_code, 200)
            html = response.get_data(as_text=True)
            self.assertNotIn("已调整", html, "未调整时不应渲染已调整徽章")
            self.assertIn('name="auto_lodging" value="300.0"', html)
            self.assertIn('name="lodging" type="number" step="0.01" min="0" value="300.0"', html,
                          "住宿单元格应显示来源金额 300")

    # ---- 8. 已调整时页面显示手改值 + 徽章，且隐藏域仍回传来源金额 ----
    def test_rendered_input_shows_manual_value_with_badge(self):
        self._recalc()
        with self.module.app.app_context():
            db = self.module.db()
            db.execute("update customer_reimbursement_items set lodging = 250 where customer_reimbursement_id = ?",
                       (self.reimbursement_id,))
            db.commit()
        with self.module.app.test_client() as client:
            with client.session_transaction() as session:
                session["user_id"] = self.user_id
            response = client.get(f"/service-orders/{self.order_id}/customer-reimbursement")
            self.assertEqual(response.status_code, 200)
            html = response.get_data(as_text=True)
            self.assertIn("已调整", html, "手改过应渲染已调整徽章")
            self.assertIn('value="300.0"', html, "隐藏域仍应回传来源金额")
            self.assertIn('name="lodging" type="number" step="0.01" min="0" value="250.0"', html,
                          "住宿单元格应显示手改后的 250")

    # ---- 9. manual_review（选择报销来源）路径同样保持手改值 ----
    def test_manual_review_mode_preserves_override(self):
        with self.module.app.app_context():
            db = self.module.db()
            db.execute("update customer_reimbursements set expense_selection_mode = 'manual_review' where id = ?",
                       (self.reimbursement_id,))
            db.execute(
                """
                insert into customer_reimbursement_expense_links (
                    customer_reimbursement_id, expense_item_id, amount_snapshot,
                    project_snapshot, expense_status_snapshot, selected_by, selected_at
                ) values (?, (select id from expense_items limit 1), 300, 'Accommodation/Lodging',
                          'approved', ?, '2026-09-15T12:00:00')
                """,
                (self.reimbursement_id, self.user_id),
            )
            db.commit()
        totals = self._recalc()
        self.assertEqual(totals["lodging_total"], 300, "manual_review 下来源金额应生效")
        item = self._item()
        self.assertEqual(item["auto_lodging"], 300)
        item["lodging"] = 250.0
        item.pop("auto_expense_sources", None)
        self._recalc([item])
        after = self._item()
        self.assertEqual(after["lodging"], 250.0, "manual_review 下保存后手改值不应被冲掉")
        self.assertEqual(self._effective(after, "lodging"), 250.0)


class ReimbursementPlanARegressionTest(unittest.TestCase):
    """方案 A 回归：翻倍根因修复、MRO 口径、gate 三出口、Excel 合计行。

    翻倍根因②（同人异日拆行）+ 根因③（PG 日期类型）由
    _reimbursement_date_key / _fallback_row_for_worker 修复。
    """

    @classmethod
    def setUpClass(cls):
        cls.temp_dir = tempfile.TemporaryDirectory()
        module_path = Path(cls.temp_dir.name) / "app_plan_a.py"
        shutil.copyfile(REPO_DIR / "app.py", module_path)
        spec = importlib.util.spec_from_file_location("invoice_tool_plan_a_app", module_path)
        cls.module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.module)
        cls.module.app.config.update(TESTING=True, SECRET_KEY="test-secret")

    @classmethod
    def tearDownClass(cls):
        cls.temp_dir.cleanup()

    def setUp(self):
        with self.module.app.app_context():
            db = self.module.db()
            # 关闭 FK 检查，避免删除顺序（或上轮遗留孤立行）触发约束失败。
            db.execute("PRAGMA foreign_keys=OFF")
            so_row = db.execute("select id from service_orders where order_number='SO-PLANA'").fetchone()
            so_id = so_row[0] if so_row else None
            if so_id is not None:
                db.execute(
                    "delete from customer_reimbursement_items where customer_reimbursement_id in "
                    "(select id from customer_reimbursements where service_order_id=?)",
                    (so_id,),
                )
                db.execute("delete from customer_reimbursements where service_order_id=?", (so_id,))
                db.execute(
                    "delete from expense_items where expense_id in (select id from expenses where service_order_id=?)",
                    (so_id,),
                )
                db.execute("delete from expenses where service_order_id=?", (so_id,))
            db.execute("delete from users where email='plan-a-admin@example.com'")
            db.execute("delete from service_orders where order_number='SO-PLANA'")
            lodging_project_id = db.execute("select id from projects where name='Accommodation/Lodging'").fetchone()
            if lodging_project_id is None:
                lodging_project_id = db.execute(
                    "insert into projects (name,name_key,project_type,unit_price,is_active,created_at) "
                    "values ('Accommodation/Lodging','accommodation/lodging','expense',0,1,'2026-09-15T12:00:00')"
                ).lastrowid
            else:
                lodging_project_id = lodging_project_id[0]
            self.lodging_project_id = lodging_project_id
            user_id = db.execute(
                "insert into users (name,email,password_hash,role,created_at) values "
                "('Plan A Admin','plan-a-admin@example.com','x','admin','2026-09-15T12:00:00')"
            ).lastrowid
            order_id = db.execute(
                "insert into service_orders (order_number,client_name,site_address,client_order_number,"
                "start_date,created_by,created_at) values ('SO-PLANA','C','addr','CL','2026-09-15',?, "
                "'2026-09-15T12:00:00')",
                (user_id,),
            ).lastrowid
            reimb_id = db.execute(
                "insert into customer_reimbursements (service_order_id,file_name,stored_filename,status,"
                "expense_selection_mode,created_by,created_at) values (?, 's.pdf','s.pdf','draft','legacy',?, "
                "'2026-09-15T12:00:00')",
                (order_id, user_id),
            ).lastrowid
            db.execute("PRAGMA foreign_keys=ON")
            db.commit()
            self.user_id = user_id
            self.order_id = order_id
            self.reimb_id = reimb_id

    def _make_expense(self, expense_date, amount, status="approved", project="Accommodation/Lodging"):
        with self.module.app.app_context():
            db = self.module.db()
            eid = db.execute(
                "insert into expenses (service_order_id,expense_number,project,expense_date,amount,status,"
                "created_by,beneficiary_id,created_at,updated_at) values (?,?,?,?,?,?,?,?, '2026-09-15T12:00:00',"
                "'2026-09-15T12:00:00')",
                (self.order_id, f"EXP-{expense_date}-{amount}-{status}", project, expense_date, amount,
                 status, self.user_id, self.user_id),
            ).lastrowid
            line_key = f"line-{eid}-0"
            db.execute(
                "insert into expense_items (expense_id,project_id,project,amount,sort_order,line_key) "
                "values (?,?,?,?,0,?)",
                (eid, self.lodging_project_id, project, amount, line_key),
            )
            db.commit()
            return eid, line_key

    def _seed_item(self, worker, project_date, with_source=None):
        with self.module.app.app_context():
            db = self.module.db()
            sources = json.dumps({"lodging": [with_source]}) if with_source else "{}"
            db.execute(
                "insert into customer_reimbursement_items (customer_reimbursement_id,worker_name,project_date,"
                "auto_expense_sources,sort_order) values (?,?,?,?,0)",
                (self.reimb_id, worker, project_date, sources),
            )
            db.commit()

    def _merge_result(self, disable_fallback=False):
        with self.module.app.app_context():
            rows = self.module.customer_reimbursement_items(self.reimb_id)
            original = None
            if disable_fallback:
                original = self.module._fallback_row_for_worker
                self.module._fallback_row_for_worker = lambda cw, wk: None
            result = self.module.merge_approved_expenses_into_customer_reimbursement(
                rows, self.order_id, reimbursement_id=self.reimb_id
            )
            if disable_fallback and original is not None:
                self.module._fallback_row_for_worker = original
            return result

    def _count_worker(self, result, name):
        return sum(1 for r in result if r.get("worker_name") == name)

    def _reimb_row(self):
        with self.module.app.app_context():
            return self.module.db().execute(
                "select * from customer_reimbursements where id=?", (self.reimb_id,)
            ).fetchone()

    # ---- 日期双栈归一 ----
    def test_date_key_normalizes_variants(self):
        f = self.module._reimbursement_date_key
        from datetime import date, datetime
        self.assertEqual(f(date(2026, 9, 18)), "2026-09-18")
        self.assertEqual(f(datetime(2026, 9, 18, 0, 0, 0)), "2026-09-18")
        self.assertEqual(f("2026-09-18 00:00:00"), "2026-09-18")
        self.assertEqual(f("2026-09-18T00:00:00"), "2026-09-18")
        self.assertEqual(f("2026-09-18"), "2026-09-18")
        self.assertEqual(f(None), "")
        self.assertEqual(f(""), "")

    # ---- 翻倍回归：禁用 fallback 应拆成 2 行（复现缺陷）----
    def test_double_count_regression_without_fallback(self):
        self._seed_item("Plan A Admin", "2026-09-15")
        self._make_expense("2026-09-20", 300, status="approved")
        result = self._merge_result(disable_fallback=True)
        self.assertEqual(self._count_worker(result, "Plan A Admin"), 2)

    # ---- 修复后：同人异日回落到已有行，不新增 ----
    def test_no_double_count_with_fallback(self):
        self._seed_item("Plan A Admin", "2026-09-15")
        self._make_expense("2026-09-20", 300, status="approved")
        result = self._merge_result()
        self.assertEqual(self._count_worker(result, "Plan A Admin"), 1)
        row = [r for r in result if r.get("worker_name") == "Plan A Admin"][0]
        self.assertEqual(float(row.get("auto_lodging") or 0), 300)
        self.assertEqual(float(row.get("lodging") or 0), 0)

    # ---- 未调整（manual==auto）入库归一为 0 ----
    def test_untouched_normalized_to_zero_on_save(self):
        self._seed_item("Plan A Admin", "2026-09-15")
        self._make_expense("2026-09-20", 300, status="approved")
        result = self._merge_result()
        with self.module.app.app_context():
            self.module.save_customer_reimbursement_items(self.reimb_id, result)
            self.module.db().commit()
            item = dict(self.module.customer_reimbursement_items(self.reimb_id)[0])
        self.assertEqual(float(item["auto_lodging"]), 300)
        self.assertEqual(float(item["lodging"]), 0, "manual==auto 应归一为 0")

    # ---- MRO 参数绝不能改变合计；total 自洽 ----
    def test_totals_mro_param_ignored_and_self_consistent(self):
        items = [
            {"labor_total": 100, "total": 570.03, "lodging": 0, "auto_lodging": 50, "airfare": 0,
             "auto_airfare": 0, "baggage": 0, "auto_baggage": 0, "rental_car": 0, "auto_rental_car": 0,
             "fuel": 0, "auto_fuel": 0, "parking": 0, "auto_parking": 0, "taxi": 0, "auto_taxi": 0,
             "mileage_total": 50, "other": 0, "auto_other": 370.03},
            {"labor_total": 50, "total": 100, "lodging": 0, "auto_lodging": 20, "airfare": 0,
             "auto_airfare": 0, "baggage": 0, "auto_baggage": 0, "rental_car": 0, "auto_rental_car": 0,
             "fuel": 0, "auto_fuel": 0, "parking": 0, "auto_parking": 0, "taxi": 0, "auto_taxi": 0,
             "mileage_total": 30, "other": 0, "auto_other": 0},
        ]
        a = self.module.customer_reimbursement_totals(items, mro_supplies_total=0, rental_fuel_total=155.5)
        b = self.module.customer_reimbursement_totals(items, mro_supplies_total=370.03, rental_fuel_total=155.5)
        self.assertEqual(a["total_amount"], b["total_amount"], "MRO 参数绝不能改变合计")
        self.assertEqual(a["mro_supplies_total"], 0)
        self.assertEqual(b["mro_supplies_total"], 370.03)
        diff = a["total_amount"] - (a["labor_total"] + a["lodging_total"] + a["other_total"] + a["mileage_total"])
        self.assertAlmostEqual(diff, 155.5, places=4)

    def _gate(self, allow_pending=False):
        # gate 内部调 db()/money()，必须在 app context 内执行。
        with self.module.app.app_context():
            return self.module.customer_reimbursement_gate_error(self._reimb_row(), allow_pending=allow_pending)

    # ---- gate 三出口 ----
    def test_gate_includable_hard_block(self):
        self._make_expense("2026-09-20", 300, status="approved")
        gate = self._gate()
        self.assertIsNotNone(gate)
        self.assertEqual(gate[0], "pending_sources")

    def test_gate_pending_soft_block_default(self):
        self._make_expense("2026-09-20", 300, status="pending")
        gate = self._gate()
        self.assertIsNotNone(gate)
        self.assertEqual(gate[0], "pending_approval")

    def test_gate_pending_allowed_with_confirm(self):
        self._make_expense("2026-09-20", 300, status="pending")
        gate = self._gate(allow_pending=True)
        self.assertIsNone(gate)

    def test_gate_returned_not_blocked(self):
        self._make_expense("2026-09-20", 300, status="returned")
        gate = self._gate()
        self.assertIsNone(gate)

    def test_gate_all_included_none(self):
        eid, line_key = self._make_expense("2026-09-20", 300, status="approved")
        self._seed_item("Plan A Admin", "2026-09-15", with_source={"expense_id": eid, "line_key": line_key, "amount": 300})
        gate = self._gate()
        self.assertIsNone(gate)

    # ---- Excel 合计行 ----
    def test_excel_totals_row_present(self):
        with self.module.app.app_context():
            db = self.module.db()
            db.execute(
                "insert into customer_reimbursement_items (customer_reimbursement_id,worker_name,project_date,"
                "standard_hours,standard_rate,labor_total,total,miles,mileage_rate,mileage_total,sort_order) "
                "values (?, 'A','2026-09-15',2,50,100,300,10,0.5,5,0)",
                (self.reimb_id,),
            )
            db.execute(
                "insert into customer_reimbursement_items (customer_reimbursement_id,worker_name,project_date,"
                "standard_hours,standard_rate,labor_total,total,miles,mileage_rate,mileage_total,sort_order) "
                "values (?, 'B','2026-09-16',3,50,150,450,20,0.5,10,1)",
                (self.reimb_id,),
            )
            db.execute(
                "update customer_reimbursements set total_amount=750, rental_fuel_total=0 where id=?",
                (self.reimb_id,),
            )
            db.commit()
        with self.module.app.test_client() as client:
            with client.session_transaction() as s:
                s["user_id"] = self.user_id
            resp = client.get(f"/customer-reimbursements/{self.reimb_id}/download.xlsx")
            self.assertEqual(resp.status_code, 200)
            from io import BytesIO
            from openpyxl import load_workbook
            wb = load_workbook(BytesIO(resp.get_data()))
            rows = list(wb.active.iter_rows(values_only=True))
            self.assertEqual(len(rows), 5, "表头+2明细+合计+说明")
            total_row = rows[-2]
            self.assertEqual(total_row[0], "合计")
            self.assertAlmostEqual(float(total_row[-1]), 750, places=2)

    # ---- 路由级集成：submit 被 gate 拦 → sync 计入 → 再 submit 成功 ----
    def test_route_submit_gate_then_sync_then_submit(self):
        self._make_expense("2026-09-20", 300, status="approved")
        self._seed_item("Plan A Admin", "2026-09-15")
        with self.module.app.test_client() as client:
            with client.session_transaction() as s:
                s["user_id"] = self.user_id
            # 1) 已审核报销未计入 → submit 被 gate 拦截，保持 draft
            resp = client.post(
                f"/service-orders/{self.order_id}/customer-reimbursement",
                data={"action": "submit"}, follow_redirects=False,
            )
            self.assertEqual(resp.status_code, 302)
            with self.module.app.app_context():
                status = self.module.db().execute(
                    "select status from customer_reimbursements where id=?", (self.reimb_id,)
                ).fetchone()["status"]
                self.assertEqual(status, "draft", "gate 拦截后应保持 draft")
            # 2) sync-sources 计入并重算快照
            resp2 = client.post(
                f"/service-orders/{self.order_id}/customer-reimbursement/sync-sources",
                data={}, follow_redirects=False,
            )
            self.assertEqual(resp2.status_code, 302)
            with self.module.app.app_context():
                item = dict(self.module.customer_reimbursement_items(self.reimb_id)[0])
                self.assertAlmostEqual(float(item["auto_lodging"]), 300, "已审核报销应计入")
                reimb = self.module.db().execute(
                    "select total_amount from customer_reimbursements where id=?", (self.reimb_id,)
                ).fetchone()
                self.assertAlmostEqual(float(reimb["total_amount"]), 300, places=2)
            # 3) 再 submit 无 includable → 成功进入 submitted（需回传完整表单）
            resp3 = client.post(
                f"/service-orders/{self.order_id}/customer-reimbursement",
                data={
                    "action": "submit",
                    "worker_name": ["Plan A Admin"],
                    "project_date": ["2026-09-15"],
                    "standard_hours": ["0"],
                    "transport_hours": ["0"],
                    "public_transport_hours": ["0"],
                    "overtime_hours": ["0"],
                    "holiday_hours": ["0"],
                    "lodging": ["0"], "auto_lodging": ["300"],
                    "airfare": ["0"], "auto_airfare": ["0"],
                    "baggage": ["0"], "auto_baggage": ["0"],
                    "rental_car": ["0"], "auto_rental_car": ["0"],
                    "fuel": ["0"], "auto_fuel": ["0"],
                    "parking": ["0"], "auto_parking": ["0"],
                    "taxi": ["0"], "auto_taxi": ["0"],
                    "miles": ["0"],
                    "other": ["0"], "auto_other": ["0"],
                    "source_report_id": [""],
                    "source_worker_user_id": [""],
                },
                follow_redirects=False,
            )
            self.assertEqual(resp3.status_code, 302)
            with self.module.app.app_context():
                status = self.module.db().execute(
                    "select status from customer_reimbursements where id=?", (self.reimb_id,)
                ).fetchone()["status"]
                self.assertEqual(status, "submitted")

    # ---- 结算页横幅：已审核未计入时渲染「计入并重算」按钮 ----
    def test_pending_sources_banner_renders(self):
        self._make_expense("2026-09-20", 300, status="approved")
        with self.module.app.test_client() as client:
            with client.session_transaction() as s:
                s["user_id"] = self.user_id
            resp = client.get(f"/service-orders/{self.order_id}/customer-reimbursement")
            self.assertEqual(resp.status_code, 200)
            html = resp.get_data(as_text=True)
            self.assertIn("已审核报销尚未计入", html)
            self.assertIn("计入并重算", html)
            self.assertIn("sync-sources", html)


if __name__ == "__main__":
    unittest.main()
