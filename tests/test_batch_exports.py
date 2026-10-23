from io import BytesIO
from openpyxl import load_workbook
from invoice_tool.batch_exports.service import build_master_ledger, settlement_rows


def expense_amount(item, key):
    return item.get(key, 0)


def test_rental_fuel_is_visible_detail_row():
    reimbursement = {"rental_fuel_total": 25.5, "total_amount": 125.5}
    items = [{
        "worker_name": "Ethan", "project_date": "2026-10-20",
        "standard_hours": 1, "transport_hours": 0, "public_transport_hours": 0,
        "overtime_hours": 0, "holiday_hours": 0, "labor_total": 100,
        "lodging": 0, "airfare": 0, "baggage": 0, "rental_car": 0,
        "fuel": 0, "parking": 0, "taxi": 0, "mileage_total": 0,
        "other": 0, "total": 100,
    }]
    rows = settlement_rows(reimbursement, items, expense_amount)
    assert len(rows) == 2
    assert rows[1][1] == "租车油费"
    assert rows[1][13] == 25.5
    assert rows[1][18] == 25.5


def test_master_ledger_matches_reference_columns_and_fuel():
    reimbursement = {"rental_fuel_total": 25.5, "total_amount": 125.5}
    order = {"order_number": "WO-1001"}
    items = [{
        "worker_name": "Ethan", "project_date": "2026-10-20",
        "standard_hours": 1, "transport_hours": 2, "overtime_hours": 0, "holiday_hours": 0,
        "standard_rate": 70, "transport_rate": 35, "overtime_rate": 105, "holiday_rate": 105,
        "labor_total": 140, "lodging": 0, "airfare": 0, "baggage": 0, "rental_car": 0,
        "fuel": 0, "parking": 0, "taxi": 0, "miles": 10, "mileage_rate": 1,
        "mileage_total": 10, "other": 0, "total": 150,
    }]
    output = build_master_ledger([(reimbursement, order, items)], expense_amount=expense_amount)
    wb = load_workbook(BytesIO(output.getvalue()), data_only=True)
    ws = wb["工单结算"]
    assert ws.max_column == 25
    assert ws["A1"].value == "派工单号"
    assert ws["M2"].value == "工时费\n合计"
    assert ws["N2"].value == "住宿"
    assert ws["O2"].value == "机票"
    assert ws["R2"].value == "加油"
    assert ws["A3"].value == "WO-1001"
    assert any(rng.min_col == 1 and rng.max_col == 1 and rng.min_row == 3 and rng.max_row > 3 for rng in ws.merged_cells.ranges)
    assert ws["R4"].value == 25.5
    assert ws["Y4"].value == 25.5
    assert "租车油费已在" in ws.cell(ws.max_row, 1).value


def test_snapshot_discount_is_visible_adjustment():
    reimbursement = {"rental_fuel_total": 20, "total_amount": 108}
    items = [{
        "worker_name": "Ethan", "project_date": "2026-10-20",
        "standard_hours": 0, "transport_hours": 0, "public_transport_hours": 0,
        "overtime_hours": 0, "holiday_hours": 0, "labor_total": 100,
        "lodging": 0, "airfare": 0, "baggage": 0, "rental_car": 0,
        "fuel": 0, "parking": 0, "taxi": 0, "mileage_total": 0,
        "other": 0, "total": 100,
    }]
    rows = settlement_rows(reimbursement, items, expense_amount)
    assert rows[-1][1] == "结算折扣/调整"
    assert rows[-1][18] == -12
    assert round(sum(row[18] for row in rows), 2) == 108



def test_mileage_proofs_export_is_grouped_by_order_and_report(tmp_path):
    import sqlite3
    import zipfile
    from contextlib import nullcontext
    from types import SimpleNamespace

    from invoice_tool.batch_exports.service import EXPORT_LABELS, run_export_job

    conn = sqlite3.connect(":memory:")
    conn.row_factory = sqlite3.Row
    conn.executescript("""
        create table batch_export_jobs (
            id text primary key, user_id integer, export_type text, order_ids text,
            status text, progress integer, created_at text, started_at text,
            completed_at text, file_name text, stored_path text, error text
        );
        create table owners (id integer primary key, name text);
        create table buyers (id integer primary key, owner text, owner_id integer);
        create table service_orders (
            id integer primary key, order_number text, client_order_number text, buyer_id integer
        );
        create table service_reports (
            id integer primary key, service_order_id integer, report_date text
        );
        create table service_report_attachments (
            id integer primary key, report_id integer, category text,
            original_filename text, stored_filename text
        );
    """)
    conn.execute("insert into service_orders(id, order_number) values(1, 'WO-1001')")
    conn.execute("insert into service_reports(id, service_order_id, report_date) values(10, 1, '2026-10-20')")
    conn.execute("insert into service_reports(id, service_order_id, report_date) values(11, 1, '2026-10-21')")
    conn.execute(
        "insert into service_report_attachments(id, report_id, category, original_filename, stored_filename) values(100,10,'mileage_proof','route.png','stored-route.png')"
    )
    conn.execute(
        "insert into batch_export_jobs(id,user_id,export_type,order_ids,status,progress,created_at) values('job-1',1,'mileage_proofs','[1]','queued',0,'2026-10-23')"
    )
    conn.commit()

    proof = tmp_path / "route.png"
    proof.write_bytes(b"mileage-proof")
    app = SimpleNamespace(app_context=lambda: nullcontext(), logger=SimpleNamespace(exception=lambda *a, **k: None))

    run_export_job(
        "job-1",
        app=app,
        db=lambda: conn,
        data_dir=tmp_path,
        build_service_report_docx=lambda *a, **k: None,
        build_simple_xlsx=lambda *a, **k: None,
        customer_reimbursement_items=lambda *a, **k: [],
        expense_amount=lambda *a, **k: 0,
        render_invoice_pdf=lambda *a, **k: None,
        report_attachment_path=lambda attachment: str(proof) if attachment["id"] == 100 else None,
    )

    job = conn.execute("select * from batch_export_jobs where id='job-1'").fetchone()
    assert job["status"] == "completed"
    assert job["progress"] == 100
    assert job["file_name"].startswith(EXPORT_LABELS["mileage_proofs"] + "_")
    with zipfile.ZipFile(job["stored_path"], "r") as archive:
        names = archive.namelist()
        assert "\u6279\u5bfc\u91cc\u7a0b/WO-1001/2026-10-20/route.png" in names
        assert "\u7f3a\u5931\u8bf4\u660e.txt" in names
        missing = archive.read("\u7f3a\u5931\u8bf4\u660e.txt").decode("utf-8")
        assert "WO-1001 / 2026-10-21: \u65e0\u91cc\u7a0b\u4f50\u8bc1" in missing
